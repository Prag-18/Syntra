import io
import json
import re
import logging
import httpx
import pdfplumber
import docx

from core.config import GEMINI_API_KEY, GEMINI_API_URL, MAX_TOKENS

logger = logging.getLogger(__name__)

class ResumeParseError(Exception):
    """Exception raised when text extraction or structure parsing fails."""
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.message = message
        self.status_code = status_code

class GeminiServiceError(Exception):
    """Exception raised when Gemini LLM service fails or times out."""
    def __init__(self, message: str, status_code: int = 502):
        super().__init__(message)
        self.message = message
        self.status_code = status_code

def sanitize_text(text: str) -> str:
    """Strip control characters from extracted text while keeping newlines and tabs."""
    if not text:
        return ""
    # Strip ASCII control characters except \t and \n
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract raw text from a PDF file using pdfplumber."""
    try:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        raw_text = "\n".join(text_parts)
        cleaned = sanitize_text(raw_text)
        if not cleaned.strip():
            raise ResumeParseError(
                "Scanned or image-only PDF detected with no extractable text. Please fill the candidate form manually.",
                status_code=422
            )
        return cleaned
    except ResumeParseError:
        raise
    except Exception as e:
        logger.error(f"PDF extraction error: {type(e).__name__}")
        raise ResumeParseError(f"Failed to extract text from PDF: {str(e)}", status_code=422)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract raw text from a DOCX file using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = [p.text for p in doc.paragraphs if p.text]
        # Also check tables in docx
        for table in doc.tables:
            for row in table.rows:
                row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_parts.append(row_text)
        raw_text = "\n".join(text_parts)
        cleaned = sanitize_text(raw_text)
        if not cleaned.strip():
            raise ResumeParseError(
                "DOCX file contains no extractable text. Please fill the candidate form manually.",
                status_code=422
            )
        return cleaned
    except ResumeParseError:
        raise
    except Exception as e:
        logger.error(f"DOCX extraction error: {type(e).__name__}")
        raise ResumeParseError(f"Failed to extract text from DOCX: {str(e)}", status_code=422)

async def extract_candidate_from_resume(raw_text: str) -> dict:
    """
    Calls Gemini model via HTTP REST API to extract structured candidate information from resume text.
    Score fields (system_design_score, coding_score, communication_score) are explicitly excluded
    from the LLM schema and defaulted to None with needs_assessment=True.
    """
    prompt = f"""
    You are an expert recruiter AI. Analyze the following resume text and extract candidate profile details.
    Return ONLY a valid JSON object matching this exact schema (no markdown, no prose wrapper):
    {{
      "full_name": "string",
      "current_role": "string",
      "company": "string",
      "years_experience": integer,
      "skills": ["string"],
      "bio": "string",
      "mentoring_signals": "string",
      "communication_signals": "string",
      "collaboration_signals": "string",
      "public_presence": "string",
      "referral_notes": "string"
    }}

    Rules:
    1. Extract full_name, current_role, company, years_experience, skills, bio summary.
    2. Extract public_presence links (e.g. GitHub or LinkedIn URLs if present in resume).
    3. Do NOT invent or hallucinate data not found in the resume.
    4. Return 0 for years_experience if unknown or not mentioned.
    5. Return empty strings or empty arrays for missing fields.

    Resume Text:
    {raw_text}
    """

    params = {"key": GEMINI_API_KEY}
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "maxOutputTokens": MAX_TOKENS,
            "responseMimeType": "application/json"
        }
    }

    try:
        async with httpx.AsyncClient(timeout=30.0) as client:
            res = await client.post(GEMINI_API_URL, params=params, json=payload)
            res.raise_for_status()
            data = res.json()

        raw_llm_text = data["candidates"][0]["content"]["parts"][0]["text"].strip()
        if raw_llm_text.startswith("```json"):
            raw_llm_text = raw_llm_text[7:]
        elif raw_llm_text.startswith("```"):
            raw_llm_text = raw_llm_text[3:]
        if raw_llm_text.endswith("```"):
            raw_llm_text = raw_llm_text[:-3]

        parsed = json.loads(raw_llm_text.strip())
    except (httpx.TimeoutException, httpx.HTTPStatusError, httpx.RequestError) as e:
        logger.error(f"Gemini API request failed during resume extraction: {type(e).__name__}")
        raise GeminiServiceError(f"Gemini extraction service unavailable: {str(e)}", status_code=502)
    except (KeyError, IndexError, json.JSONDecodeError) as e:
        logger.error(f"Failed to parse Gemini response during resume extraction: {type(e).__name__}")
        raise GeminiServiceError("Failed to parse candidate data from Gemini response.", status_code=502)

    # Sanitize and assemble final candidate dictionary with explicit null score defaults and needs_assessment=True
    skills_raw = parsed.get("skills", [])
    if isinstance(skills_raw, str):
        skills_list = [s.strip() for s in skills_raw.split(",") if s.strip()]
    elif isinstance(skills_raw, list):
        skills_list = [str(s).strip() for s in skills_raw if str(s).strip()]
    else:
        skills_list = []

    try:
        yoe = int(parsed.get("years_experience", 0))
    except (ValueError, TypeError):
        yoe = 0

    return {
        "full_name": str(parsed.get("full_name") or "Unknown Candidate").strip(),
        "current_role": str(parsed.get("current_role") or "Candidate").strip(),
        "company": str(parsed.get("company") or "Independent").strip(),
        "years_experience": yoe,
        "skills": skills_list,
        "bio": str(parsed.get("bio") or "").strip(),
        "mentoring_signals": str(parsed.get("mentoring_signals") or "").strip(),
        "communication_signals": str(parsed.get("communication_signals") or "").strip(),
        "collaboration_signals": str(parsed.get("collaboration_signals") or "").strip(),
        "public_presence": str(parsed.get("public_presence") or "").strip(),
        "referral_notes": str(parsed.get("referral_notes") or "").strip(),
        # SCORE FIELDS — EXPLICIT DEFAULT, NEVER HALLUCINATE:
        "system_design_score": None,
        "coding_score": None,
        "communication_score": None,
        "needs_assessment": True
    }
